#!/usr/bin/env python3
from __future__ import annotations

import base64
import json
import re
import shutil
import sqlite3
import threading
import uuid
import warnings
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
DEFAULT_DB_PATH = PROJECT_DIR / "veri" / "apartman.db"
DEFAULT_IMPORT_XLSX = PROJECT_DIR / "kaynak" / "ilk_veri.xlsx"
DEFAULT_EXPORT_DIR = PROJECT_DIR / "raporlar"
DOCUMENTS_DIR = PROJECT_DIR / "belgeler"
SCHEMA_VERSION = 3

CHARGE_TYPES = ("aidat", "demirbas")
CHARGE_LABELS = {"aidat": "Aidat", "demirbas": "Demirbaş"}
EXPENSE_CATEGORIES = [
    "Asansor",
    "Bahce/Peyzaj",
    "Bakim",
    "Banka",
    "Elektrik",
    "Gorevli",
    "Kirtasiye",
    "Onarim",
    "Temizlik",
    "Yonetim",
    "Ilaclama",
]
BUDGET_TYPES = ["Aidat", "Demirbaş"]
RESIDENT_TYPES = ["Ev Sahibi", "Kiraci"]
DOCUMENT_TYPES = ["Makbuz", "Fatura", "Toplanti Tutanağı", "Karar Defteri", "Denetci Raporu", "Diger"]
PERIOD_PATTERN = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
MONTH_NAMES = {
    "ocak": 1,
    "subat": 2,
    "mart": 3,
    "nisan": 4,
    "mayis": 5,
    "haziran": 6,
    "temmuz": 7,
    "agustos": 8,
    "eylul": 9,
    "ekim": 10,
    "kasim": 11,
    "aralik": 12,
}
MONTH_LABELS = {
    1: "Ocak",
    2: "Şubat",
    3: "Mart",
    4: "Nisan",
    5: "Mayıs",
    6: "Haziran",
    7: "Temmuz",
    8: "Ağustos",
    9: "Eylül",
    10: "Ekim",
    11: "Kasım",
    12: "Aralık",
}


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def normalize_text(value: Any) -> str:
    return (
        str(value or "")
        .strip()
        .lower()
        .replace("\u0307", "")
        .replace("ı", "i")
        .replace("ğ", "g")
        .replace("ş", "s")
        .replace("ö", "o")
        .replace("ü", "u")
        .replace("ç", "c")
    )


def normalize_charge_type(value: str) -> str:
    normalized = normalize_text(value)
    if normalized not in CHARGE_TYPES:
        raise ValueError("Yukumluluk turu gecersiz.")
    return normalized


def normalize_budget_type(value: str) -> str:
    normalized = normalize_text(value)
    if normalized not in CHARGE_TYPES:
        raise ValueError("Butce turu gecersiz.")
    return CHARGE_LABELS[normalized]


def to_cents(value: Any) -> int:
    if value in (None, ""):
        return 0
    try:
        amount = Decimal(str(value).replace("₺", "").replace("TL", "").strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Gecersiz tutar: {value}") from exc
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def cents_to_number(value: Any) -> float:
    return round(int(value or 0) / 100, 2)


def money_text(amount_cents: int) -> str:
    amount = Decimal(int(amount_cents or 0)) / Decimal("100")
    text = f"{amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if text.endswith(",00"):
        text = text[:-3]
    return f"{text} TL"


def safe_number(value: Any) -> float:
    return cents_to_number(to_cents(value))


def parse_period_key(value: str) -> tuple[int, int]:
    match = PERIOD_PATTERN.match(str(value or "").strip())
    if not match:
        raise ValueError("Donem YYYY-AA formatinda olmalidir. Ornek: 2026-09")
    return int(match.group(1)), int(match.group(2))


def period_label(period_key: str) -> str:
    year, month = parse_period_key(period_key)
    return f"{MONTH_LABELS[month]} {year}"


def current_period_key(today: date | None = None) -> str:
    today = today or date.today()
    return f"{today.year:04d}-{today.month:02d}"


def period_is_overdue(period_key: str, as_of: date | None = None) -> bool:
    return parse_period_key(period_key) < parse_period_key(current_period_key(as_of))


def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if digits.startswith("90"):
        return digits
    if digits.startswith("0"):
        return "90" + digits[1:]
    if digits.startswith("5") and len(digits) == 10:
        return "90" + digits
    return digits


def parse_amount_from_text(value: Any) -> Decimal | None:
    text = str(value or "")
    match = re.search(r"[-+]?\d[\d.,]*", text)
    if not match:
        return None
    raw = match.group(0)
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def month_number(value: Any) -> int | None:
    normalized = normalize_text(value)
    for name, number in MONTH_NAMES.items():
        if name in normalized:
            return number
    return None


def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
    return bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?", (table,)).fetchone())


def _column_exists(conn: sqlite3.Connection, table: str, column: str) -> bool:
    return any(row[1] == column for row in conn.execute(f"PRAGMA table_info({table})").fetchall())


class ApartmentRepository:
    def __init__(self, db_path: Path = DEFAULT_DB_PATH):
        self.db_path = Path(db_path)
        self.lock = threading.RLock()

    def connect(self) -> sqlite3.Connection:
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("PRAGMA busy_timeout = 5000")
        return conn

    def initialize(self) -> None:
        with self.lock:
            existed = self.db_path.exists()
            with self.connect() as conn:
                self._create_base_schema(conn)
                self._create_v3_schema(conn)
                version = int(self._meta(conn, "schema_version") or "1")
            if version < SCHEMA_VERSION:
                if existed:
                    backup_path = self._backup_database("v3_gecis_oncesi")
                else:
                    backup_path = None
                with self.connect() as conn:
                    self._migrate_to_v3(conn, backup_path)

    def _create_base_schema(self, conn: sqlite3.Connection) -> None:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS app_meta (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS units (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unit_number INTEGER NOT NULL UNIQUE,
                resident_name TEXT NOT NULL,
                resident_type TEXT NOT NULL,
                phone TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS monthly_obligations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unit_id INTEGER NOT NULL REFERENCES units(id) ON DELETE CASCADE,
                month_key TEXT NOT NULL,
                month_label TEXT NOT NULL,
                period_order INTEGER NOT NULL,
                charge_type TEXT NOT NULL,
                due_amount REAL NOT NULL DEFAULT 0,
                paid_amount REAL NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(unit_id, month_key, charge_type)
            );
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                expense_date TEXT NOT NULL,
                document_no TEXT,
                vendor TEXT NOT NULL,
                description TEXT NOT NULL,
                category TEXT NOT NULL,
                budget_type TEXT NOT NULL,
                amount REAL NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS activity_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_type TEXT NOT NULL,
                entity_type TEXT NOT NULL,
                entity_id TEXT,
                payload_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            """
        )
        conn.execute("INSERT OR IGNORE INTO app_meta(key, value) VALUES ('schema_version', '1')")

    def _create_v3_schema(self, conn: sqlite3.Connection) -> None:
        for table, column, definition in (
            ("units", "aidat_muaf", "INTEGER NOT NULL DEFAULT 0"),
            ("units", "demirbas_muaf", "INTEGER NOT NULL DEFAULT 0"),
            ("expenses", "amount_cents", "INTEGER NOT NULL DEFAULT 0"),
            ("expenses", "document_id", "INTEGER"),
            ("activity_log", "action", "TEXT NOT NULL DEFAULT ''"),
            ("activity_log", "old_value", "TEXT"),
            ("activity_log", "new_value", "TEXT"),
            ("activity_log", "actor", "TEXT NOT NULL DEFAULT 'yonetici'"),
        ):
            if not _column_exists(conn, table, column):
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")

        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS schema_migrations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                version INTEGER NOT NULL UNIQUE,
                applied_at TEXT NOT NULL,
                notes TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS periods (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                period_key TEXT NOT NULL UNIQUE CHECK(period_key GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]'),
                year INTEGER NOT NULL,
                month INTEGER NOT NULL CHECK(month BETWEEN 1 AND 12),
                label TEXT NOT NULL,
                starts_on TEXT NOT NULL,
                ends_on TEXT NOT NULL,
                is_closed INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS annual_budgets (
                year INTEGER PRIMARY KEY,
                aidat_cents INTEGER NOT NULL DEFAULT 0 CHECK(aidat_cents >= 0),
                demirbas_cents INTEGER NOT NULL DEFAULT 0 CHECK(demirbas_cents >= 0),
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS obligations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unit_id INTEGER NOT NULL REFERENCES units(id) ON DELETE CASCADE,
                period_id INTEGER NOT NULL REFERENCES periods(id) ON DELETE CASCADE,
                charge_type TEXT NOT NULL CHECK(charge_type IN ('aidat', 'demirbas')),
                accrual_cents INTEGER NOT NULL DEFAULT 0 CHECK(accrual_cents >= 0),
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(unit_id, period_id, charge_type)
            );
            CREATE TABLE IF NOT EXISTS payment_transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unit_id INTEGER NOT NULL REFERENCES units(id) ON DELETE CASCADE,
                payment_date TEXT NOT NULL,
                amount_cents INTEGER NOT NULL CHECK(amount_cents > 0),
                payment_method TEXT NOT NULL DEFAULT 'banka',
                receipt_no TEXT NOT NULL UNIQUE,
                note TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS payment_allocations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                payment_id INTEGER NOT NULL REFERENCES payment_transactions(id) ON DELETE CASCADE,
                obligation_id INTEGER NOT NULL REFERENCES obligations(id) ON DELETE CASCADE,
                amount_cents INTEGER NOT NULL CHECK(amount_cents > 0),
                UNIQUE(payment_id, obligation_id)
            );
            CREATE TABLE IF NOT EXISTS account_movements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                movement_date TEXT NOT NULL,
                unit_id INTEGER REFERENCES units(id) ON DELETE SET NULL,
                period_id INTEGER REFERENCES periods(id) ON DELETE SET NULL,
                obligation_id INTEGER REFERENCES obligations(id) ON DELETE SET NULL,
                payment_id INTEGER REFERENCES payment_transactions(id) ON DELETE SET NULL,
                expense_id INTEGER REFERENCES expenses(id) ON DELETE SET NULL,
                charge_type TEXT CHECK(charge_type IN ('aidat', 'demirbas')),
                movement_type TEXT NOT NULL CHECK(movement_type IN ('tahakkuk', 'tahsilat', 'gider', 'duzeltme')),
                direction TEXT NOT NULL CHECK(direction IN ('borc', 'alacak')),
                amount_cents INTEGER NOT NULL CHECK(amount_cents >= 0),
                description TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS late_fee_settings (
                id INTEGER PRIMARY KEY CHECK(id = 1),
                enabled INTEGER NOT NULL DEFAULT 0,
                rate_basis_points INTEGER NOT NULL DEFAULT 0,
                calculation_method TEXT NOT NULL DEFAULT 'simple_monthly',
                effective_date TEXT,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS late_fee_transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                obligation_id INTEGER NOT NULL REFERENCES obligations(id) ON DELETE CASCADE,
                amount_cents INTEGER NOT NULL CHECK(amount_cents >= 0),
                calculated_at TEXT NOT NULL,
                calculation_method TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS board_decisions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                decision_no TEXT NOT NULL UNIQUE,
                decision_date TEXT NOT NULL,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_type TEXT NOT NULL,
                file_name TEXT NOT NULL,
                storage_key TEXT NOT NULL UNIQUE,
                mime_type TEXT NOT NULL DEFAULT 'application/octet-stream',
                file_size INTEGER NOT NULL DEFAULT 0,
                unit_id INTEGER REFERENCES units(id) ON DELETE SET NULL,
                decision_id INTEGER REFERENCES board_decisions(id) ON DELETE SET NULL,
                uploaded_at TEXT NOT NULL,
                uploaded_by TEXT NOT NULL DEFAULT 'yonetici'
            );
            CREATE TABLE IF NOT EXISTS audit_reports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                report_date TEXT NOT NULL,
                report_type TEXT NOT NULL,
                payload_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            """
        )
        conn.execute(
            "INSERT OR IGNORE INTO late_fee_settings(id, enabled, rate_basis_points, calculation_method, updated_at) VALUES (1, 0, 0, 'simple_monthly', ?)",
            (now_text(),),
        )
        conn.execute("UPDATE expenses SET amount_cents = CAST(ROUND(amount * 100, 0) AS INTEGER) WHERE amount_cents = 0 AND amount <> 0")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_obligations_period ON obligations(period_id, charge_type)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_obligations_unit ON obligations(unit_id, period_id)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_payments_unit_date ON payment_transactions(unit_id, payment_date)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_movements_date ON account_movements(movement_date, id)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_expenses_category ON expenses(category, budget_type)")

    def _meta(self, conn: sqlite3.Connection, key: str) -> str | None:
        row = conn.execute("SELECT value FROM app_meta WHERE key = ?", (key,)).fetchone()
        return row["value"] if row else None

    def _set_meta(self, conn: sqlite3.Connection, key: str, value: str) -> None:
        conn.execute("INSERT OR REPLACE INTO app_meta(key, value) VALUES (?, ?)", (key, value))

    def _backup_database(self, prefix: str) -> Path | None:
        if not self.db_path.exists():
            return None
        backup_dir = self.db_path.parent / "yedekler"
        backup_dir.mkdir(parents=True, exist_ok=True)
        target = backup_dir / f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
        with self.connect() as conn:
            conn.execute("PRAGMA wal_checkpoint(FULL)")
        shutil.copy2(self.db_path, target)
        return target

    def _migrate_to_v3(self, conn: sqlite3.Connection, backup_path: Path | None) -> None:
        source_year = self._infer_legacy_year(conn)
        legacy_rows = conn.execute("SELECT * FROM monthly_obligations ORDER BY id").fetchall()
        budget_values: dict[int, dict[str, int]] = {}
        for row in legacy_rows:
            order = int(row["period_order"] or 1)
            legacy_month = normalize_text(row["month_key"])
            month = MONTH_NAMES.get(legacy_month, max(1, min(12, order if order <= 12 else 1)))
            period_key = str(row["month_key"] or "")
            if not PERIOD_PATTERN.match(period_key):
                period_key = f"{source_year:04d}-{month:02d}"
            year, month = parse_period_key(period_key)
            period_id = self._ensure_period_conn(conn, period_key, year, month)
            charge_type = normalize_charge_type(row["charge_type"])
            due_cents = to_cents(row["due_amount"])
            paid_cents = to_cents(row["paid_amount"])
            budget_values.setdefault(year, {}).setdefault(charge_type, due_cents)
            conn.execute(
                "INSERT OR IGNORE INTO obligations(unit_id, period_id, charge_type, accrual_cents, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
                (row["unit_id"], period_id, charge_type, due_cents, now_text(), now_text()),
            )
            obligation = conn.execute(
                "SELECT id FROM obligations WHERE unit_id = ? AND period_id = ? AND charge_type = ?",
                (row["unit_id"], period_id, charge_type),
            ).fetchone()
            period_row = conn.execute("SELECT period_key, starts_on FROM periods WHERE id = ?", (period_id,)).fetchone()
            if obligation and not conn.execute("SELECT 1 FROM account_movements WHERE obligation_id = ? AND movement_type = 'tahakkuk'", (obligation["id"],)).fetchone():
                self._add_movement_conn(conn, period_row["starts_on"], row["unit_id"], period_id, obligation["id"], None, None, charge_type, "tahakkuk", "borc", due_cents, f"{period_row['period_key']} {CHARGE_LABELS[charge_type]} tahakkuku")
            if obligation and paid_cents > 0:
                receipt = f"MIG-{int(row['id']):08d}"
                conn.execute(
                    "INSERT OR IGNORE INTO payment_transactions(unit_id, payment_date, amount_cents, payment_method, receipt_no, note, created_at, updated_at) VALUES (?, ?, ?, 'migrasyon', ?, 'Eski tablodan aktarildi', ?, ?)",
                    (row["unit_id"], f"{year:04d}-{month:02d}-01", paid_cents, receipt, now_text(), now_text()),
                )
                payment = conn.execute("SELECT id FROM payment_transactions WHERE receipt_no = ?", (receipt,)).fetchone()
                conn.execute(
                    "INSERT OR IGNORE INTO payment_allocations(payment_id, obligation_id, amount_cents) VALUES (?, ?, ?)",
                    (payment["id"], obligation["id"], paid_cents),
                )
                if not conn.execute("SELECT 1 FROM account_movements WHERE payment_id = ?", (payment["id"],)).fetchone():
                    self._add_movement_conn(conn, period_row["starts_on"], row["unit_id"], period_id, obligation["id"], payment["id"], None, charge_type, "tahsilat", "alacak", paid_cents, f"{period_row['period_key']} {CHARGE_LABELS[charge_type]} aktarimi")
        for year, values in budget_values.items():
            conn.execute(
                "INSERT OR IGNORE INTO annual_budgets(year, aidat_cents, demirbas_cents, created_at, updated_at) VALUES (?, ?, ?, ?, ?)",
                (year, values.get("aidat", 0), values.get("demirbas", 0), now_text(), now_text()),
            )
        for expense in conn.execute("SELECT id, expense_date, description, budget_type, amount_cents, amount FROM expenses").fetchall():
            if not conn.execute("SELECT 1 FROM account_movements WHERE expense_id = ?", (expense["id"],)).fetchone():
                amount_cents = int(expense["amount_cents"] or to_cents(expense["amount"]))
                self._add_movement_conn(
                    conn,
                    expense["expense_date"],
                    None,
                    None,
                    None,
                    None,
                    expense["id"],
                    normalize_text(expense["budget_type"]),
                    "gider",
                    "borc",
                    amount_cents,
                    expense["description"],
                )
        conn.execute("INSERT OR IGNORE INTO schema_migrations(version, applied_at, notes) VALUES (3, ?, ?)", (now_text(), "Eski aylik yukumlulukler donem bazli modele aktarildi"))
        self._set_meta(conn, "schema_version", str(SCHEMA_VERSION))
        self._set_meta(conn, "v3_migration_backup", str(backup_path or ""))
        self.log_activity(conn, "migration_v3", "database", None, {"backup": str(backup_path or ""), "legacy_rows": len(legacy_rows)}, actor="system")

    def _infer_legacy_year(self, conn: sqlite3.Connection) -> int:
        source = self._meta(conn, "import_source_excel") or ""
        match = re.search(r"(20\d{2})", source)
        return int(match.group(1)) if match else 2026

    def _ensure_period_conn(self, conn: sqlite3.Connection, period_key: str, year: int | None = None, month: int | None = None) -> int:
        parsed_year, parsed_month = parse_period_key(period_key)
        year = year or parsed_year
        month = month or parsed_month
        starts_on = f"{year:04d}-{month:02d}-01"
        if month == 12:
            next_month = date(year + 1, 1, 1)
        else:
            next_month = date(year, month + 1, 1)
        ends_on = str(date.fromordinal(next_month.toordinal() - 1))
        conn.execute(
            "INSERT OR IGNORE INTO periods(period_key, year, month, label, starts_on, ends_on, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (period_key, year, month, period_label(period_key), starts_on, ends_on, now_text(), now_text()),
        )
        return conn.execute("SELECT id FROM periods WHERE period_key = ?", (period_key,)).fetchone()["id"]

    def log_activity(
        self,
        conn: sqlite3.Connection,
        action: str,
        entity_type: str,
        entity_id: str | int | None,
        payload: dict[str, Any] | None = None,
        old_value: Any = None,
        new_value: Any = None,
        actor: str = "yonetici",
    ) -> None:
        payload = payload or {}
        text_payload = json.dumps(payload, ensure_ascii=False, default=str)
        conn.execute(
            "INSERT INTO activity_log(event_type, entity_type, entity_id, payload_json, created_at, action, old_value, new_value, actor) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (action, entity_type, str(entity_id) if entity_id is not None else None, text_payload, now_text(), action, json.dumps(old_value, ensure_ascii=False, default=str) if old_value is not None else None, json.dumps(new_value, ensure_ascii=False, default=str) if new_value is not None else None, actor),
        )

    def has_seed_data(self) -> bool:
        with self.connect() as conn:
            return bool(conn.execute("SELECT COUNT(*) AS count FROM units WHERE active = 1").fetchone()["count"])

    def import_from_excel(self, excel_path: Path, replace_existing: bool = False) -> dict:
        if not excel_path.exists():
            raise FileNotFoundError(f"Excel dosyasi bulunamadi: {excel_path}")
        workbook = load_workbook(excel_path, data_only=False)
        if "Gelir_Takibi" not in workbook.sheetnames or "Gider_Takibi" not in workbook.sheetnames:
            raise ValueError("Excel dosyasinda Gelir_Takibi ve Gider_Takibi sayfalari bulunmali.")
        source = self._read_excel_source(workbook)
        imported_units = 0
        imported_obligations = 0
        imported_expenses = 0
        with self.lock, self.connect() as conn:
            if replace_existing:
                conn.executescript(
                    """
                    DELETE FROM payment_allocations;
                    DELETE FROM payment_transactions;
                    DELETE FROM account_movements;
                    DELETE FROM obligations;
                    DELETE FROM periods;
                    DELETE FROM annual_budgets;
                    DELETE FROM expenses;
                    DELETE FROM monthly_obligations;
                    """
                )
            now = now_text()
            source_numbers = set()
            for item in source["units"]:
                unit_number = int(item["unit_number"])
                source_numbers.add(unit_number)
                existing = conn.execute("SELECT id FROM units WHERE unit_number = ?", (unit_number,)).fetchone()
                if existing:
                    conn.execute("UPDATE units SET resident_name = ?, resident_type = ?, phone = ?, active = 1, updated_at = ? WHERE id = ?", (item["resident_name"], item["resident_type"], item["phone"], now, existing["id"]))
                    unit_id = existing["id"]
                else:
                    unit_id = conn.execute("INSERT INTO units(unit_number, resident_name, resident_type, phone, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)", (unit_number, item["resident_name"], item["resident_type"], item["phone"], now, now)).lastrowid
                imported_units += 1
                for obligation in item["obligations"]:
                    if self._unit_is_exempt(conn, unit_id, obligation["charge_type"]):
                        continue
                    period_id = self._ensure_period_conn(conn, obligation["period_key"])
                    self._upsert_budget_value(conn, int(obligation["period_key"][:4]), obligation["charge_type"], obligation["accrual_cents"])
                    obligation_id = self._create_obligation_conn(conn, unit_id, period_id, obligation["charge_type"], obligation["accrual_cents"], now)
                    imported_obligations += 1
                    if obligation["paid_cents"] > 0:
                        payment_id = conn.execute("INSERT INTO payment_transactions(unit_id, payment_date, amount_cents, payment_method, receipt_no, note, created_at, updated_at) VALUES (?, ?, ?, 'excel', ?, 'Excel aktarimi', ?, ?)", (unit_id, f"{obligation['period_key']}-01", obligation["paid_cents"], f"EXL-{uuid.uuid4().hex[:12].upper()}", now, now)).lastrowid
                        conn.execute("INSERT INTO payment_allocations(payment_id, obligation_id, amount_cents) VALUES (?, ?, ?)", (payment_id, obligation_id, obligation["paid_cents"]))
                        self._add_movement_conn(conn, f"{obligation['period_key']}-01", unit_id, period_id, obligation_id, payment_id, None, obligation["charge_type"], "tahsilat", "alacak", obligation["paid_cents"], "Excel aktarimi")
            if source_numbers:
                marks = ",".join("?" for _ in source_numbers)
                conn.execute(f"UPDATE units SET active = 0, updated_at = ? WHERE unit_number NOT IN ({marks})", (now, *sorted(source_numbers)))
            for expense in source["expenses"]:
                conn.execute("INSERT INTO expenses(expense_date, document_no, vendor, description, category, budget_type, amount, amount_cents, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", (expense["expense_date"], expense["document_no"], expense["vendor"], expense["description"], expense["category"], expense["budget_type"], cents_to_number(expense["amount_cents"]), expense["amount_cents"], now, now))
                expense_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                self._add_movement_conn(conn, expense["expense_date"], None, None, None, None, expense_id, normalize_text(expense["budget_type"]), "gider", "borc", expense["amount_cents"], expense["description"])
                imported_expenses += 1
            self._set_meta(conn, "import_source_excel", str(excel_path))
            self._set_meta(conn, "last_imported_at", now)
            self.log_activity(conn, "excel_import", "database", None, {"excel_path": str(excel_path), "replace_existing": replace_existing, "units": imported_units, "obligations": imported_obligations, "expenses": imported_expenses})
        return {"units": imported_units, "obligations": imported_obligations, "expenses": imported_expenses, "excel_path": str(excel_path)}

    def _read_excel_source(self, workbook) -> dict[str, Any]:
        gelir = workbook["Gelir_Takibi"]
        gider = workbook["Gider_Takibi"]
        source_year = self._year_from_sheet(gelir)
        aidat_fields: list[tuple[int, int, int, int]] = []
        demirbas_fields: list[tuple[int, int, int]] = []
        for col in range(1, gelir.max_column + 1):
            header = gelir.cell(3, col).value
            detail = gelir.cell(4, col).value
            month = month_number(header)
            if month and "aidat" in normalize_text(detail):
                amount = parse_amount_from_text(detail) or Decimal("0")
                aidat_fields.append((month, col, col + 1, to_cents(amount)))
        demirbas_start = next((col for col in range(1, gelir.max_column + 1) if "demirbas" in normalize_text(gelir.cell(3, col).value)), None)
        if demirbas_start:
            for col in range(demirbas_start, gelir.max_column + 1):
                detail = gelir.cell(4, col).value
                month = month_number(detail)
                if month:
                    amount = parse_amount_from_text(detail) or Decimal("0")
                    demirbas_fields.append((month, col, to_cents(amount)))
        units = []
        for row in range(5, gelir.max_row + 1):
            if gelir.cell(row, 1).value in (None, ""):
                continue
            unit = {"unit_number": int(gelir.cell(row, 1).value), "resident_name": str(gelir.cell(row, 2).value or "").strip(), "resident_type": str(gelir.cell(row, 3).value or "Ev Sahibi").strip(), "phone": str(gelir.cell(row, 4).value or "").strip(), "obligations": []}
            for month, due_col, paid_col, default_cents in aidat_fields:
                due = to_cents(gelir.cell(row, due_col).value or default_cents / 100)
                paid = to_cents(gelir.cell(row, paid_col).value)
                unit["obligations"].append({"period_key": f"{source_year:04d}-{month:02d}", "charge_type": "aidat", "accrual_cents": due, "paid_cents": paid})
            for month, paid_col, default_cents in demirbas_fields:
                paid = to_cents(gelir.cell(row, paid_col).value)
                unit["obligations"].append({"period_key": f"{source_year:04d}-{month:02d}", "charge_type": "demirbas", "accrual_cents": default_cents, "paid_cents": paid})
            units.append(unit)
        expenses = []
        total_row = self._find_expense_total_row(gider)
        for row in range(4, total_row):
            amount = gider.cell(row, 8).value
            if amount in (None, ""):
                continue
            raw_date = gider.cell(row, 2).value
            expense_date = raw_date.strftime("%Y-%m-%d") if hasattr(raw_date, "strftime") else str(raw_date)
            budget_type = normalize_budget_type(str(gider.cell(row, 7).value or "Aidat"))
            expenses.append({"expense_date": expense_date, "document_no": str(gider.cell(row, 3).value or "").strip(), "vendor": str(gider.cell(row, 4).value or "").strip(), "description": str(gider.cell(row, 5).value or "").strip(), "category": str(gider.cell(row, 6).value or "Diger").strip(), "budget_type": budget_type, "amount_cents": to_cents(amount)})
        return {"units": units, "expenses": expenses}

    @staticmethod
    def _year_from_sheet(sheet) -> int:
        for row in sheet.iter_rows(min_row=1, max_row=3, values_only=True):
            match = re.search(r"(20\d{2})", " ".join(str(value or "") for value in row))
            if match:
                return int(match.group(1))
        return date.today().year

    @staticmethod
    def _find_expense_total_row(sheet) -> int:
        for row in range(4, sheet.max_row + 2):
            if "toplam gider" in normalize_text(sheet.cell(row, 7).value):
                return row
        return sheet.max_row + 1

    def _unit_is_exempt(self, conn: sqlite3.Connection, unit_id: int, charge_type: str) -> bool:
        column = "aidat_muaf" if charge_type == "aidat" else "demirbas_muaf"
        return bool(conn.execute(f"SELECT {column} FROM units WHERE id = ?", (unit_id,)).fetchone()[0])

    def _upsert_budget_value(self, conn: sqlite3.Connection, year: int, charge_type: str, cents: int) -> None:
        column = "aidat_cents" if charge_type == "aidat" else "demirbas_cents"
        conn.execute("INSERT OR IGNORE INTO annual_budgets(year, aidat_cents, demirbas_cents, created_at, updated_at) VALUES (?, 0, 0, ?, ?)", (year, now_text(), now_text()))
        conn.execute(f"UPDATE annual_budgets SET {column} = ?, updated_at = ? WHERE year = ?", (cents, now_text(), year))

    def _create_obligation_conn(self, conn: sqlite3.Connection, unit_id: int, period_id: int, charge_type: str, accrual_cents: int, timestamp: str | None = None) -> int:
        charge_type = normalize_charge_type(charge_type)
        if self._unit_is_exempt(conn, unit_id, charge_type):
            raise ValueError("Muaf daire icin yukumluluk olusturulamaz.")
        timestamp = timestamp or now_text()
        conn.execute("INSERT OR IGNORE INTO obligations(unit_id, period_id, charge_type, accrual_cents, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)", (unit_id, period_id, charge_type, accrual_cents, timestamp, timestamp))
        row = conn.execute("SELECT id FROM obligations WHERE unit_id = ? AND period_id = ? AND charge_type = ?", (unit_id, period_id, charge_type)).fetchone()
        if not conn.execute("SELECT 1 FROM account_movements WHERE obligation_id = ? AND movement_type = 'tahakkuk'", (row["id"],)).fetchone():
            period = conn.execute("SELECT starts_on FROM periods WHERE id = ?", (period_id,)).fetchone()
            self._add_movement_conn(conn, period["starts_on"], unit_id, period_id, row["id"], None, None, charge_type, "tahakkuk", "borc", accrual_cents, f"{period_label(conn.execute('SELECT period_key FROM periods WHERE id = ?', (period_id,)).fetchone()['period_key'])} {CHARGE_LABELS[charge_type]} tahakkuku")
        return row["id"]

    def _add_movement_conn(self, conn: sqlite3.Connection, movement_date: str, unit_id: int | None, period_id: int | None, obligation_id: int | None, payment_id: int | None, expense_id: int | None, charge_type: str | None, movement_type: str, direction: str, amount_cents: int, description: str) -> None:
        conn.execute("INSERT INTO account_movements(movement_date, unit_id, period_id, obligation_id, payment_id, expense_id, charge_type, movement_type, direction, amount_cents, description, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", (movement_date, unit_id, period_id, obligation_id, payment_id, expense_id, charge_type if charge_type in CHARGE_TYPES else None, movement_type, direction, amount_cents, description, now_text()))

    def _obligation_rows(self, conn: sqlite3.Connection, unit_id: int | None = None, overdue_only: bool = False) -> list[dict[str, Any]]:
        query = """
            SELECT o.id AS obligation_id, o.unit_id, o.period_id, o.charge_type, o.accrual_cents,
                   p.period_key, p.label AS period_label, p.starts_on,
                   u.unit_number, u.resident_name, u.resident_type, u.phone,
                   u.aidat_muaf, u.demirbas_muaf,
                   COALESCE((SELECT SUM(amount_cents) FROM payment_allocations WHERE obligation_id = o.id), 0) AS paid_cents
            FROM obligations o
            JOIN periods p ON p.id = o.period_id
            JOIN units u ON u.id = o.unit_id
            WHERE u.active = 1
        """
        params: list[Any] = []
        if unit_id is not None:
            query += " AND o.unit_id = ?"
            params.append(unit_id)
        if overdue_only:
            query += " AND p.period_key < ? AND o.accrual_cents > COALESCE((SELECT SUM(amount_cents) FROM payment_allocations WHERE obligation_id = o.id), 0)"
            params.append(current_period_key())
        query += " ORDER BY p.period_key, o.charge_type, u.unit_number"
        rows = []
        for row in conn.execute(query, params).fetchall():
            item = dict(row)
            item["remaining_cents"] = max(0, int(item["accrual_cents"]) - int(item["paid_cents"]))
            item["is_overdue"] = period_is_overdue(item["period_key"]) and item["remaining_cents"] > 0
            rows.append(item)
        return rows

    def _period_summaries(self, conn: sqlite3.Connection) -> list[dict[str, Any]]:
        rows = conn.execute("SELECT * FROM periods ORDER BY period_key").fetchall()
        result = []
        for period in rows:
            item = {"key": period["period_key"], "label": period["label"], "status": "gecmis" if period["period_key"] < current_period_key() else ("aktif" if period["period_key"] == current_period_key() else "gelecek")}
            for charge_type in CHARGE_TYPES:
                values = conn.execute("SELECT COALESCE(SUM(o.accrual_cents), 0) AS accrual, COALESCE((SELECT SUM(pa.amount_cents) FROM payment_allocations pa JOIN obligations ox ON ox.id = pa.obligation_id WHERE ox.period_id = o.period_id AND ox.charge_type = o.charge_type), 0) AS paid FROM obligations o WHERE o.period_id = ? AND o.charge_type = ?", (period["id"], charge_type)).fetchone()
                accrual = int(values["accrual"] or 0)
                paid = int(values["paid"] or 0)
                item[charge_type] = {"accrual": cents_to_number(accrual), "paid": cents_to_number(paid), "remaining": cents_to_number(max(0, accrual - paid)), "rate": round(paid / accrual, 4) if accrual else 0}
            result.append(item)
        return result

    def _expense_rows(self, conn: sqlite3.Connection, limit: int | None = 20) -> list[dict[str, Any]]:
        query = "SELECT id, expense_date, document_no, vendor, description, category, budget_type, amount_cents, document_id FROM expenses ORDER BY expense_date DESC, id DESC"
        params: tuple[Any, ...] = ()
        if limit:
            query += " LIMIT ?"
            params = (limit,)
        return [{**dict(row), "amount": cents_to_number(row["amount_cents"])} for row in conn.execute(query, params).fetchall()]

    def _expense_groups(self, conn: sqlite3.Connection) -> list[dict[str, Any]]:
        groups: dict[tuple[str, str], dict[str, Any]] = {}
        for row in self._expense_rows(conn, None):
            key = (row["category"], row["budget_type"])
            groups.setdefault(key, {"category": row["category"], "budget_type": row["budget_type"], "total": 0, "items": []})
            groups[key]["total"] += row["amount"]
            groups[key]["items"].append(row)
        return sorted(groups.values(), key=lambda item: item["total"], reverse=True)

    def _charge_summary(self, conn: sqlite3.Connection, charge_type: str) -> dict[str, Any]:
        row = conn.execute("SELECT COALESCE(SUM(o.accrual_cents), 0) AS accrual, COALESCE((SELECT SUM(pa.amount_cents) FROM payment_allocations pa JOIN obligations ox ON ox.id = pa.obligation_id WHERE ox.charge_type = ?), 0) AS paid FROM obligations o WHERE o.charge_type = ?", (charge_type, charge_type)).fetchone()
        accrual = int(row["accrual"] or 0)
        paid = int(row["paid"] or 0)
        overdue = self._obligation_rows(conn, overdue_only=True)
        overdue = [item for item in overdue if item["charge_type"] == charge_type]
        expenses = conn.execute("SELECT COALESCE(SUM(amount_cents), 0) AS total FROM expenses WHERE lower(replace(budget_type, 'ş', 's')) = ?", (charge_type,)).fetchone()["total"] or 0
        return {"accrual": cents_to_number(accrual), "paid": cents_to_number(paid), "remaining": cents_to_number(max(0, accrual - paid)), "overdue": cents_to_number(sum(item["remaining_cents"] for item in overdue)), "expense": cents_to_number(expenses), "net": cents_to_number(paid - int(expenses)), "rate": round(paid / accrual, 4) if accrual else 0, "overdue_count": len(overdue)}

    def dashboard(self) -> dict[str, Any]:
        self.initialize()
        with self.lock, self.connect() as conn:
            aidat = self._charge_summary(conn, "aidat")
            demirbas = self._charge_summary(conn, "demirbas")
            overdue_rows = self._obligation_rows(conn, overdue_only=True)
            grouped: dict[int, dict[str, Any]] = {}
            for row in overdue_rows:
                item = grouped.setdefault(row["unit_id"], {"unit_id": row["unit_id"], "daire": row["unit_number"], "isim": row["resident_name"], "durum": row["resident_type"], "telefon": row["phone"], "aidat_periodleri": [], "demirbas_periodleri": [], "aidat_detaylari": [], "demirbas_detaylari": [], "aidat_remaining": 0, "demirbas_remaining": 0, "overdue_period_count": 0})
                item["overdue_period_count"] += 1
                item[f"{row['charge_type']}_remaining"] += cents_to_number(row["remaining_cents"])
                item[f"{row['charge_type']}_periodleri"].append(row["period_key"])
                item[f"{row['charge_type']}_detaylari"].append({"donem": period_label(row["period_key"]), "tutar": cents_to_number(row["remaining_cents"]), "tutar_metni": money_text(row["remaining_cents"])})
            debtors = []
            for item in sorted(grouped.values(), key=lambda value: value["aidat_remaining"] + value["demirbas_remaining"], reverse=True):
                detaylar = item["aidat_detaylari"] + item["demirbas_detaylari"]
                mesaj_satirlari = ["Sayın Komşum, iyi günler dilerim.", "", "📌 Borç Bilgileri"]
                mesaj_satirlari.extend(f"{detay['donem']} {CHARGE_LABELS[tur]}: {detay['tutar_metni']}" for tur in CHARGE_TYPES for detay in item[f"{tur}_detaylari"])
                item["total_remaining"] = round(item["aidat_remaining"] + item["demirbas_remaining"], 2)
                mesaj_satirlari.extend(["", f"➡️ Toplam Borç: {money_text(sum(int(round(detay['tutar'] * 100)) for detay in detaylar))}"])
                item["message"] = "\n".join(mesaj_satirlari)
                item["eksik_donemler"] = [f"{detay['donem']} {CHARGE_LABELS[tur]}: {detay['tutar_metni']}" for tur in CHARGE_TYPES for detay in item[f"{tur}_detaylari"]]
                debtors.append(item)
            units = []
            for row in conn.execute("SELECT id, unit_number AS daire, resident_name AS isim, resident_type AS durum, phone, aidat_muaf, demirbas_muaf FROM units WHERE active = 1 ORDER BY unit_number").fetchall():
                units.append({**dict(row), "aidat_muaf": bool(row["aidat_muaf"]), "demirbas_muaf": bool(row["demirbas_muaf"])})
            total_collected = aidat["paid"] + demirbas["paid"]
            total_expenses = aidat["expense"] + demirbas["expense"]
            budget_rows = [dict(row) for row in conn.execute("SELECT year, aidat_cents, demirbas_cents, updated_at FROM annual_budgets ORDER BY year DESC").fetchall()]
            eligible_aidat = conn.execute("SELECT COUNT(*) FROM units WHERE active = 1 AND aidat_muaf = 0").fetchone()[0]
            eligible_demirbas = conn.execute("SELECT COUNT(*) FROM units WHERE active = 1 AND demirbas_muaf = 0").fetchone()[0]
            for row in budget_rows:
                row["aidat"] = cents_to_number(row.pop("aidat_cents")); row["demirbas"] = cents_to_number(row.pop("demirbas_cents"))
                row["aidat_beklenen_yillik"] = round(row["aidat"] * eligible_aidat * 12, 2)
                row["demirbas_beklenen_yillik"] = round(row["demirbas"] * eligible_demirbas * 12, 2)
            expense_months = []
            for row in conn.execute("SELECT substr(expense_date, 1, 7) AS period_key, budget_type, SUM(amount_cents) AS amount_cents FROM expenses GROUP BY period_key, budget_type ORDER BY period_key DESC").fetchall():
                expense_months.append({"period_key": row["period_key"], "budget_type": normalize_text(row["budget_type"]), "amount": cents_to_number(row["amount_cents"])})
            return {
                "summary": {"current_period": current_period_key(), "total_collected": round(total_collected, 2), "total_expenses": round(total_expenses, 2), "net_balance": round(total_collected - total_expenses, 2), "total_outstanding": round(aidat["remaining"] + demirbas["remaining"], 2), "debtor_count": len(debtors), "overdue_period_count": sum(item["overdue_period_count"] for item in debtors), "updated_at": now_text()},
                "aidat": aidat,
                "demirbas": demirbas,
                "periods": self._period_summaries(conn),
                "current_period": current_period_key(),
                "debtors": debtors,
                "apartments": units,
                "recent_expenses": self._expense_rows(conn),
                "expense_months": expense_months,
                "expense_groups": self._expense_groups(conn),
                "tracking": {"aidat_risk": aidat["overdue"], "demirbas_risk": demirbas["overdue"], "missing_phone_count": sum(1 for item in units if not item["phone"]), "exempt_count": sum(1 for item in units if item["aidat_muaf"] or item["demirbas_muaf"]), "overdue_period_count": sum(item["overdue_period_count"] for item in debtors)},
                "budgets": budget_rows,
                "options": {"expense_categories": EXPENSE_CATEGORIES, "budget_types": BUDGET_TYPES, "document_types": DOCUMENT_TYPES, "periods": [{"key": item["key"], "label": item["label"]} for item in self._period_summaries(conn)]},
            }

    def add_period(self, period_key: str) -> dict[str, Any]:
        year, month = parse_period_key(period_key)
        with self.lock, self.connect() as conn:
            period_id = self._ensure_period_conn(conn, period_key, year, month)
            budget = conn.execute("SELECT aidat_cents, demirbas_cents FROM annual_budgets WHERE year = ?", (year,)).fetchone()
            created = 0
            if budget:
                for unit in conn.execute("SELECT id FROM units WHERE active = 1").fetchall():
                    for charge_type, column in (("aidat", "aidat_cents"), ("demirbas", "demirbas_cents")):
                        if not self._unit_is_exempt(conn, unit["id"], charge_type):
                            self._create_obligation_conn(conn, unit["id"], period_id, charge_type, int(budget[column]))
                            created += 1
            self.log_activity(conn, "period_created", "period", period_key, {"obligations": created})
        return {"message": f"{period_label(period_key)} donemi olusturuldu.", "period_key": period_key, "obligations": created}

    def save_annual_budget(self, year: int, aidat: Any, demirbas: Any) -> dict[str, Any]:
        year = int(year)
        aidat_cents = to_cents(aidat); demirbas_cents = to_cents(demirbas)
        with self.lock, self.connect() as conn:
            conn.execute("INSERT INTO annual_budgets(year, aidat_cents, demirbas_cents, created_at, updated_at) VALUES (?, ?, ?, ?, ?) ON CONFLICT(year) DO UPDATE SET aidat_cents = excluded.aidat_cents, demirbas_cents = excluded.demirbas_cents, updated_at = excluded.updated_at", (year, aidat_cents, demirbas_cents, now_text(), now_text()))
            for period in conn.execute("SELECT id, period_key FROM periods WHERE year = ?", (year,)).fetchall():
                for unit in conn.execute("SELECT id FROM units WHERE active = 1").fetchall():
                    for charge_type, cents in (("aidat", aidat_cents), ("demirbas", demirbas_cents)):
                        if self._unit_is_exempt(conn, unit["id"], charge_type):
                            continue
                        obligation = conn.execute("SELECT id FROM obligations WHERE unit_id = ? AND period_id = ? AND charge_type = ?", (unit["id"], period["id"], charge_type)).fetchone()
                        if obligation:
                            paid = conn.execute("SELECT COALESCE(SUM(amount_cents), 0) AS paid FROM payment_allocations WHERE obligation_id = ?", (obligation["id"],)).fetchone()["paid"]
                            if not paid:
                                conn.execute("UPDATE obligations SET accrual_cents = ?, updated_at = ? WHERE id = ?", (cents, now_text(), obligation["id"]))
                        elif period["period_key"] >= current_period_key():
                            self._create_obligation_conn(conn, unit["id"], period["id"], charge_type, cents)
            self.log_activity(conn, "annual_budget_updated", "annual_budget", year, {}, new_value={"aidat": cents_to_number(aidat_cents), "demirbas": cents_to_number(demirbas_cents)})
        return {"message": f"{year} yili butcesi kaydedildi.", "year": year, "aidat": cents_to_number(aidat_cents), "demirbas": cents_to_number(demirbas_cents), "eligible_units": self._eligible_unit_counts(year)}

    def _eligible_unit_counts(self, year: int) -> dict[str, int]:
        with self.connect() as conn:
            return {"aidat": conn.execute("SELECT COUNT(*) FROM units WHERE active = 1 AND aidat_muaf = 0").fetchone()[0], "demirbas": conn.execute("SELECT COUNT(*) FROM units WHERE active = 1 AND demirbas_muaf = 0").fetchone()[0]}

    def update_unit_exemptions(self, unit_id: int, aidat_muaf: bool, demirbas_muaf: bool) -> dict[str, Any]:
        with self.lock, self.connect() as conn:
            before = dict(conn.execute("SELECT aidat_muaf, demirbas_muaf FROM units WHERE id = ?", (unit_id,)).fetchone() or {})
            if not before:
                raise ValueError("Daire bulunamadi.")
            conn.execute("UPDATE units SET aidat_muaf = ?, demirbas_muaf = ?, updated_at = ? WHERE id = ?", (int(aidat_muaf), int(demirbas_muaf), now_text(), unit_id))
            self.log_activity(conn, "unit_exemption_updated", "unit", unit_id, {}, old_value=before, new_value={"aidat_muaf": aidat_muaf, "demirbas_muaf": demirbas_muaf})
            unit = conn.execute("SELECT unit_number FROM units WHERE id = ?", (unit_id,)).fetchone()
        return {"message": f"Daire {unit['unit_number']} muafiyetleri guncellendi."}

    def add_payment(self, unit_id: int, charge_type: str, period_key: str, amount: Any, payment_date: str | None = None, payment_method: str = "banka", note: str = "") -> dict[str, Any]:
        charge_type = normalize_charge_type(charge_type)
        amount_cents = to_cents(amount)
        if amount_cents <= 0:
            raise ValueError("Tutar sifirdan buyuk olmali.")
        parse_period_key(period_key)
        payment_date = payment_date or date.today().isoformat()
        with self.lock, self.connect() as conn:
            obligation = conn.execute("SELECT o.id, o.accrual_cents, o.period_id, u.unit_number, u.resident_name FROM obligations o JOIN units u ON u.id = o.unit_id JOIN periods p ON p.id = o.period_id WHERE o.unit_id = ? AND o.charge_type = ? AND p.period_key = ?", (unit_id, charge_type, period_key)).fetchone()
            if not obligation:
                raise ValueError("Bu daire ve donem icin yukumluluk bulunamadi.")
            paid = conn.execute("SELECT COALESCE(SUM(amount_cents), 0) AS paid FROM payment_allocations WHERE obligation_id = ?", (obligation["id"],)).fetchone()["paid"]
            remaining = int(obligation["accrual_cents"]) - int(paid or 0)
            if amount_cents > remaining:
                raise ValueError("Odeme kalan yukumlulugu asiyor.")
            receipt_no = f"TAH-{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6].upper()}"
            payment_id = conn.execute("INSERT INTO payment_transactions(unit_id, payment_date, amount_cents, payment_method, receipt_no, note, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (unit_id, payment_date, amount_cents, payment_method or "banka", receipt_no, note.strip(), now_text(), now_text())).lastrowid
            conn.execute("INSERT INTO payment_allocations(payment_id, obligation_id, amount_cents) VALUES (?, ?, ?)", (payment_id, obligation["id"], amount_cents))
            self._add_movement_conn(conn, payment_date, unit_id, obligation["period_id"], obligation["id"], payment_id, None, charge_type, "tahsilat", "alacak", amount_cents, f"{period_key} {CHARGE_LABELS[charge_type]} tahsilati")
            self.log_activity(conn, "payment_added", "payment", payment_id, {"receipt_no": receipt_no, "period_key": period_key, "charge_type": charge_type, "amount": cents_to_number(amount_cents)})
            unit = conn.execute("SELECT unit_number, resident_name FROM units WHERE id = ?", (unit_id,)).fetchone()
        pdf_path = self._create_receipt_pdf(receipt_no, unit["unit_number"], unit["resident_name"], payment_date, cents_to_number(amount_cents), charge_type, period_key)
        return {"message": f"Daire {unit['unit_number']} icin tahsilat kaydedildi.", "receipt_no": receipt_no, "receipt_path": str(pdf_path) if pdf_path else None}

    def update_phone(self, unit_id: int, phone: str) -> dict[str, str]:
        clean_phone = str(phone or "").strip()
        with self.connect() as conn:
            row = conn.execute("SELECT unit_number, phone FROM units WHERE id = ?", (unit_id,)).fetchone()
            if not row:
                raise ValueError("Daire bulunamadi.")
            conn.execute("UPDATE units SET phone = ?, updated_at = ? WHERE id = ?", (clean_phone, now_text(), unit_id))
            self.log_activity(conn, "unit_phone_updated", "unit", unit_id, {"phone": clean_phone}, old_value={"phone": row["phone"]}, new_value={"phone": clean_phone})
        return {"message": f"Daire {row['unit_number']} telefon bilgisi guncellendi."}

    def _create_receipt_pdf(self, receipt_no: str, unit_number: int, resident_name: str, payment_date: str, amount: float, charge_type: str, period_key: str) -> Path | None:
        try:
            from pdf_rapor import create_payment_receipt
            return create_payment_receipt(DEFAULT_EXPORT_DIR / "makbuzlar", receipt_no, unit_number, resident_name, payment_date, amount, CHARGE_LABELS[charge_type], period_key)
        except ImportError:
            return None

    def create_debt_notice(self, unit_id: int) -> dict[str, Any]:
        with self.connect() as conn:
            unit = conn.execute("SELECT id, unit_number, resident_name, phone FROM units WHERE id = ?", (unit_id,)).fetchone()
            if not unit:
                raise ValueError("Daire bulunamadi.")
            rows = self._obligation_rows(conn, unit_id=unit_id, overdue_only=True)
        try:
            from pdf_rapor import create_debt_notice
            path = create_debt_notice(DEFAULT_EXPORT_DIR / "borc_bildirimleri", unit["unit_number"], unit["resident_name"], rows, date.today().isoformat())
        except ImportError:
            path = None
        return {"message": "Borc bildirimi hazirlandi.", "path": str(path) if path else None, "daire": unit["unit_number"]}

    def add_expense(self, expense_date: str, vendor: str, description: str, category: str, budget_type: str, amount: Any, document_no: str = "", document_id: int | None = None) -> dict[str, Any]:
        amount_cents = to_cents(amount)
        if amount_cents <= 0:
            raise ValueError("Tutar sifirdan buyuk olmali.")
        budget_type = normalize_budget_type(budget_type)
        if not category.strip():
            raise ValueError("Gider kategorisi bos olamaz.")
        datetime.strptime(expense_date, "%Y-%m-%d")
        with self.lock, self.connect() as conn:
            expense_id = conn.execute("INSERT INTO expenses(expense_date, document_no, vendor, description, category, budget_type, amount, amount_cents, document_id, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", (expense_date, document_no.strip() or None, vendor.strip(), description.strip(), category.strip(), budget_type, cents_to_number(amount_cents), amount_cents, document_id, now_text(), now_text())).lastrowid
            self._add_movement_conn(conn, expense_date, None, None, None, None, expense_id, normalize_text(budget_type), "gider", "borc", amount_cents, description.strip())
            self.log_activity(conn, "expense_added", "expense", expense_id, {"vendor": vendor, "amount": cents_to_number(amount_cents), "budget_type": budget_type})
        return {"message": f"{vendor} gideri eklendi.", "expense_id": expense_id}

    def get_account(self, unit_id: int) -> dict[str, Any]:
        with self.connect() as conn:
            unit = conn.execute("SELECT id, unit_number, resident_name, phone FROM units WHERE id = ?", (unit_id,)).fetchone()
            if not unit:
                raise ValueError("Daire bulunamadi.")
            movements = [dict(row) for row in conn.execute("SELECT am.movement_date, am.movement_type, am.direction, am.charge_type, am.amount_cents, am.description, am.period_id, p.period_key, am.obligation_id, am.payment_id FROM account_movements am LEFT JOIN periods p ON p.id = am.period_id WHERE am.unit_id = ? AND NOT (am.movement_type = 'tahakkuk' AND COALESCE((SELECT SUM(pa.amount_cents) FROM payment_allocations pa WHERE pa.obligation_id = am.obligation_id), 0) >= am.amount_cents) ORDER BY am.movement_date, am.id", (unit_id,)).fetchall()]
            for movement in movements:
                movement["amount"] = cents_to_number(movement.pop("amount_cents"))
                movement["charge_label"] = CHARGE_LABELS.get(movement.get("charge_type"), "")
        return {"unit": dict(unit), "movements": movements}

    def save_late_fee_settings(self, enabled: bool, rate: Any, calculation_method: str, effective_date: str | None) -> dict[str, Any]:
        rate_decimal = Decimal(str(rate or 0))
        rate_basis_points = int((rate_decimal * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        with self.connect() as conn:
            conn.execute("UPDATE late_fee_settings SET enabled = ?, rate_basis_points = ?, calculation_method = ?, effective_date = ?, updated_at = ? WHERE id = 1", (int(enabled), rate_basis_points, calculation_method or "simple_monthly", effective_date or None, now_text()))
            self.log_activity(conn, "late_fee_settings_updated", "late_fee_settings", 1, {}, new_value={"enabled": enabled, "rate": float(rate_decimal), "calculation_method": calculation_method})
        return {"message": "Gecikme tazminati ayarlari kaydedildi.", "enabled": bool(enabled), "rate": float(rate_decimal), "calculation_method": calculation_method}

    def get_late_fee_settings(self) -> dict[str, Any]:
        with self.connect() as conn:
            row = conn.execute("SELECT enabled, rate_basis_points, calculation_method, effective_date FROM late_fee_settings WHERE id = 1").fetchone()
        return {"enabled": bool(row["enabled"]), "rate": row["rate_basis_points"] / 100, "calculation_method": row["calculation_method"], "effective_date": row["effective_date"]}

    def get_decisions(self) -> list[dict[str, Any]]:
        with self.connect() as conn:
            return [dict(row) for row in conn.execute("SELECT * FROM board_decisions ORDER BY decision_date DESC, id DESC").fetchall()]

    def save_decision(self, decision_no: str, decision_date: str, title: str, body: str, decision_id: int | None = None) -> dict[str, Any]:
        if decision_id:
            with self.connect() as conn:
                conn.execute("UPDATE board_decisions SET decision_no = ?, decision_date = ?, title = ?, body = ?, updated_at = ? WHERE id = ?", (decision_no, decision_date, title, body, now_text(), decision_id))
                self.log_activity(conn, "board_decision_updated", "board_decision", decision_id, {"decision_no": decision_no})
            return {"message": "Karar guncellendi.", "id": decision_id}
        with self.connect() as conn:
            decision_id = conn.execute("INSERT INTO board_decisions(decision_no, decision_date, title, body, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)", (decision_no, decision_date, title, body, now_text(), now_text())).lastrowid
            self.log_activity(conn, "board_decision_added", "board_decision", decision_id, {"decision_no": decision_no})
        return {"message": "Karar kaydedildi.", "id": decision_id}

    def delete_decision(self, decision_id: int) -> dict[str, str]:
        with self.connect() as conn:
            conn.execute("DELETE FROM board_decisions WHERE id = ?", (decision_id,))
            self.log_activity(conn, "board_decision_deleted", "board_decision", decision_id)
        return {"message": "Karar silindi."}

    def save_document(self, document_type: str, file_name: str, mime_type: str, content_base64: str, unit_id: int | None = None, decision_id: int | None = None) -> dict[str, Any]:
        if document_type not in DOCUMENT_TYPES:
            raise ValueError("Belge turu gecersiz.")
        safe_name = Path(file_name).name
        if not safe_name:
            raise ValueError("Belge adi bulunamadi.")
        raw = base64.b64decode(content_base64)
        key = f"{datetime.now().strftime('%Y/%m')}/{uuid.uuid4().hex}_{safe_name}"
        target = DOCUMENTS_DIR / key
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(raw)
        with self.connect() as conn:
            document_id = conn.execute("INSERT INTO documents(document_type, file_name, storage_key, mime_type, file_size, unit_id, decision_id, uploaded_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (document_type, safe_name, key, mime_type or "application/octet-stream", len(raw), unit_id, decision_id, now_text())).lastrowid
            self.log_activity(conn, "document_added", "document", document_id, {"file_name": safe_name, "document_type": document_type})
        return {"message": "Belge arsive eklendi.", "id": document_id, "file_name": safe_name}

    def get_documents(self) -> list[dict[str, Any]]:
        with self.connect() as conn:
            return [dict(row) for row in conn.execute("SELECT id, document_type, file_name, mime_type, file_size, unit_id, decision_id, uploaded_at FROM documents ORDER BY uploaded_at DESC").fetchall()]

    def get_document_path(self, document_id: int) -> tuple[Path, str]:
        with self.connect() as conn:
            row = conn.execute("SELECT storage_key, file_name, mime_type FROM documents WHERE id = ?", (document_id,)).fetchone()
        if not row:
            raise FileNotFoundError("Belge bulunamadi.")
        path = (DOCUMENTS_DIR / row["storage_key"]).resolve()
        if DOCUMENTS_DIR.resolve() not in path.parents or not path.exists():
            raise FileNotFoundError("Belge dosyasi bulunamadi.")
        return path, row["mime_type"]

    def inspector_report(self) -> dict[str, Any]:
        data = self.dashboard()
        with self.connect() as conn:
            undocumented = conn.execute("SELECT COUNT(*) FROM expenses WHERE (document_no IS NULL OR trim(document_no) = '') AND document_id IS NULL").fetchone()[0]
            data["inspector"] = {"aidat_tahsilat_orani": data["aidat"]["rate"], "demirbas_tahsilat_orani": data["demirbas"]["rate"], "kasa_bakiyesi": data["summary"]["net_balance"], "kasa_bakiyesi_aidat": data["aidat"]["net"], "kasa_bakiyesi_demirbas": data["demirbas"]["net"], "banka_bakiyesi": None, "belgesiz_gider_sayisi": undocumented, "belgesiz_gider_kriteri": "Belge numarası girilmemiş ve arşive dosya bağlanmamış giderler.", "telefonu_eksik_kayit": data["tracking"]["missing_phone_count"], "borclu_daire_sayisi": data["summary"]["debtor_count"]}
            conn.execute("INSERT INTO audit_reports(report_date, report_type, payload_json, created_at) VALUES (?, 'denetci_ozeti', ?, ?)", (date.today().isoformat(), json.dumps(data["inspector"], ensure_ascii=False), now_text()))
        return data["inspector"]

    def list_backups(self) -> list[dict[str, Any]]:
        backup_dir = self.db_path.parent / "yedekler"
        backup_dir.mkdir(parents=True, exist_ok=True)
        return [{"name": path.name, "size": path.stat().st_size, "modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")} for path in sorted(backup_dir.glob("*.db"), key=lambda item: item.stat().st_mtime, reverse=True)]

    def create_backup(self) -> dict[str, Any]:
        target = self._backup_database("manuel")
        return {"message": "Veritabani yedeklendi.", "name": target.name if target else None}

    def restore_backup(self, backup_name: str) -> dict[str, Any]:
        if Path(backup_name).name != backup_name:
            raise ValueError("Gecersiz yedek adi.")
        candidate = self.db_path.parent / "yedekler" / backup_name
        if not candidate.exists():
            raise FileNotFoundError("Yedek dosyasi bulunamadi.")
        current_backup = self._backup_database("restore_oncesi")
        shutil.copy2(candidate, self.db_path)
        self.initialize()
        with self.connect() as conn:
            self.log_activity(conn, "database_restored", "database", None, {"backup": backup_name, "before_restore_backup": str(current_backup or "")})
        return {"message": "Veritabani geri yuklendi.", "name": backup_name}

    def export_excel_snapshot(self, export_dir: Path = DEFAULT_EXPORT_DIR) -> dict[str, Any]:
        export_dir.mkdir(parents=True, exist_ok=True)
        data = self.dashboard()
        output_path = export_dir / "guncel_rapor.xlsx"
        workbook = Workbook()
        summary = workbook.active; summary.title = "Ozet"
        units = workbook.create_sheet("Daireler")
        debts = workbook.create_sheet("Borclular")
        periods = workbook.create_sheet("Donemler")
        expenses = workbook.create_sheet("Giderler")
        summary.append(["Apartman Yonetim Uygulamasi Raporu"])
        summary.append(["Rapor Tarihi", now_text()])
        summary.append(["Aktif Donem", data["current_period"]])
        for charge_type in CHARGE_TYPES:
            item = data[charge_type]
            summary.append([CHARGE_LABELS[charge_type], "Tahakkuk", item["accrual"], "Tahsilat", item["paid"], "Bekleyen", item["remaining"], "Gecikmis", item["overdue"], "Oran", item["rate"]])
        summary.append(["Toplam Gider", data["summary"]["total_expenses"]])
        summary.append(["Net Bakiye", data["summary"]["net_balance"]])
        units.append(["Daire No", "Isim", "Durum", "Telefon", "Aidat Muaf", "Demirbas Muaf"])
        for row in data["apartments"]:
            units.append([row["daire"], row["isim"], row["durum"], row["phone"], row["aidat_muaf"], row["demirbas_muaf"]])
        debts.append(["Daire No", "Isim", "Telefon", "Aidat Eksikleri", "Demirbas Eksikleri", "Odenmeyen Donem", "Aidat Kalan", "Demirbas Kalan", "Hazir Mesaj"])
        for row in data["debtors"]:
            debts.append([row["daire"], row["isim"], row["telefon"], ", ".join(row["aidat_periodleri"]), ", ".join(row["demirbas_periodleri"]), row["overdue_period_count"], row["aidat_remaining"], row["demirbas_remaining"], row["message"]])
        periods.append(["Donem", "Durum", "Aidat Tahakkuk", "Aidat Tahsilat", "Aidat Kalan", "Demirbas Tahakkuk", "Demirbas Tahsilat", "Demirbas Kalan"])
        for row in data["periods"]:
            periods.append([row["key"], row["status"], row["aidat"]["accrual"], row["aidat"]["paid"], row["aidat"]["remaining"], row["demirbas"]["accrual"], row["demirbas"]["paid"], row["demirbas"]["remaining"]])
        expenses.append(["Tarih", "Belge No", "Firma", "Aciklama", "Kategori", "Butce", "Tutar"])
        for row in data["recent_expenses"]:
            expenses.append([row["expense_date"], row["document_no"] or "", row["vendor"], row["description"], row["category"], row["budget_type"], row["amount"]])
        for sheet in workbook.worksheets:
            for column in sheet.columns:
                width = max(len(str(cell.value or "")) for cell in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), 42)
        workbook.save(output_path)
        return {"message": "Excel raporu olusturuldu.", "path": str(output_path)}
